"""Apply a citation style to a Word (.docx) document."""
from __future__ import annotations

import copy
import os
from dataclasses import dataclass, field
from typing import Dict, Iterator, List, Optional, Tuple

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from .citations import BIBLIOGRAPHY_RE, find_citations
from .doi import looks_like_doi, resolve_doi
from .engine import BibEntry, CitationEngine
from .models import has_arabic_script
from .parsers import load_references
from .richtext import RichText


@dataclass
class ProcessReport:
    style: str
    output: str
    clusters: int = 0
    cited: List[str] = field(default_factory=list)
    missing: List[str] = field(default_factory=list)
    uncited: List[str] = field(default_factory=list)
    bibliography_entries: int = 0
    placeholder_found: bool = False

    def summary(self) -> str:
        lines = [
            f"Style:                {self.style}",
            f"Citations replaced:   {self.clusters}",
            f"Works cited:          {len(self.cited)}",
            f"Bibliography entries: {self.bibliography_entries}"
            + ("" if self.placeholder_found else "  (appended at the end: no [[BIBLIOGRAPHY]] marker)"),
            f"Output:               {self.output}",
        ]
        if self.missing:
            lines.append(f"⚠ Missing keys:       {', '.join(self.missing)}")
        if self.uncited:
            lines.append(f"ℹ Not cited:          {', '.join(self.uncited)}")
        return "\n".join(lines)


# ---------------------------------------------------------------------------- iteration
def _iter_block(el, parent) -> Iterator[Paragraph]:
    for child in el.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            for tr in child.iterchildren(qn("w:tr")):
                for tc in tr.iterchildren(qn("w:tc")):
                    yield from _iter_block(tc, parent)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_block(content, parent)


def iter_paragraphs(doc) -> Iterator[Paragraph]:
    """All body paragraphs in reading order, including tables and content controls."""
    yield from _iter_block(doc.element.body, doc._body)


# ---------------------------------------------------------------------------- run surgery
def _run_spans(p: Paragraph) -> List[Tuple[Run, int, int]]:
    spans, pos = [], 0
    for r in p.runs:
        t = r.text
        spans.append((r, pos, pos + len(t)))
        pos += len(t)
    return spans


def _split_at(p: Paragraph, offset: int) -> None:
    for run, a, b in _run_spans(p):
        if a < offset < b:
            text = run.text
            new_r = copy.deepcopy(run._r)
            run._r.addnext(new_r)
            run.text = text[: offset - a]
            Run(new_r, run._parent).text = text[offset - a:]
            return


def _apply_fmt(run: Run, fmt) -> None:
    if "i" in fmt:
        run.font.italic = True
    if "b" in fmt:
        run.font.bold = True
    if "sup" in fmt:
        run.font.superscript = True


_DROP_RPR = {qn("w:i"), qn("w:iCs"), qn("w:b"), qn("w:bCs"), qn("w:vertAlign"), qn("w:u")}


def _new_run_like(template_r, parent, text: str, fmt) -> Run:
    """A run that keeps the template's font, size, colour and direction, but not its emphasis."""
    r = copy.deepcopy(template_r)
    rPr = r.find(qn("w:rPr"))
    if rPr is not None:
        for child in list(rPr):
            if child.tag in _DROP_RPR:
                rPr.remove(child)
    run = Run(r, parent)
    run.text = text
    _apply_fmt(run, fmt)
    return run


def replace_span(p: Paragraph, start: int, end: int, rich: RichText) -> None:
    _split_at(p, start)
    _split_at(p, end)
    covered = [r for r, a, b in _run_spans(p) if a >= start and b <= end and b > a]
    if not covered:
        return
    anchor = covered[0]
    for seg in rich.segs:
        new = _new_run_like(anchor._r, anchor._parent, seg.text, seg.fmt)
        anchor._r.addprevious(new._r)
    for r in covered:
        r._r.getparent().remove(r._r)


# ---------------------------------------------------------------------------- bibliography
def _set_bidi(p: Paragraph) -> None:
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        pPr.insert(0, bidi)


def _fill_entry(para: Paragraph, entry: BibEntry, style, template_r=None) -> None:
    pf = para.paragraph_format
    if entry.label:
        indent = Cm(1.0)
        pf.left_indent, pf.first_line_indent = indent, -indent
        pf.tab_stops.add_tab_stop(indent)
    elif style.hanging_indent:
        pf.left_indent, pf.first_line_indent = Cm(1.27), Cm(-1.27)
    pf.space_after = Pt(6)
    if entry.rtl:
        _set_bidi(para)
    segs = ([(entry.label + "\t", frozenset())] if entry.label else []) + \
        [(s.text, s.fmt) for s in entry.text.segs]
    for text, fmt in segs:
        if template_r is not None:
            run = _new_run_like(template_r, para, text, fmt)
            para._p.append(run._r)
        else:
            run = para.add_run(text)
            _apply_fmt(run, fmt)
        if entry.rtl:
            run.font.rtl = True


def _insert_paragraph_after(p: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    p._p.addnext(new_p)
    para = Paragraph(new_p, p._parent)
    if p._p.pPr is not None:
        new_p.insert(0, copy.deepcopy(p._p.pPr))
    return para


def _clean_template_run(p: Paragraph):
    runs = p.runs
    if not runs:
        return None
    r = copy.deepcopy(runs[0]._r)
    for child in list(r):
        if child.tag != qn("w:rPr"):
            r.remove(child)
    return r


def _detect_doc_locale(paragraphs: List[Paragraph]) -> str:
    text = " ".join(p.text for p in paragraphs[:400])
    arabic = sum(1 for c in text if has_arabic_script(c))
    latin = sum(1 for c in text if c.isascii() and c.isalpha())
    return "fa" if arabic > latin else "en"


# ---------------------------------------------------------------------------- main entry
def process_docx(input_path, output_path, references, style="apa", locale: str = "auto",
                 include_uncited: bool = False, bib_title: Optional[str] = None,
                 persian_digits: bool = False, fetch_dois: bool = False,
                 title_case: Optional[str] = None) -> ProcessReport:
    """Replace citation markers in ``input_path`` and write the result to ``output_path``.

    ``references``: path(s) to .bib/.json/.ris files or a dict of Reference objects.
    ``bib_title``: heading for the bibliography; ``None`` = automatic, ``""`` = no heading.
    """
    library = load_references(references)
    resolver = resolve_doi if fetch_dois else None
    engine = CitationEngine(library, style, locale=locale, persian_digits=persian_digits,
                            title_case=title_case,
                            resolver=(lambda k: resolver(k) if looks_like_doi(k) else None) if resolver else None)
    doc = Document(os.fspath(input_path))
    paragraphs = list(iter_paragraphs(doc))

    # pass 1: find markers in reading order and register them (numbering, disambiguation)
    found = []
    bib_anchor: Optional[Paragraph] = None
    for p in paragraphs:
        text = "".join(r.text for r in p.runs)
        if "@" not in text and "\\cite" not in text and "[[" not in text and "{{" not in text:
            continue
        if BIBLIOGRAPHY_RE.match(text):
            bib_anchor = bib_anchor or p
            continue
        matches = find_citations(text, engine.library)
        if not matches:
            continue
        for m in matches:
            engine.register(m.cluster)
            if engine.style.superscript and not m.cluster.narrative:
                while m.start > 0 and text[m.start - 1] in " \u00a0":
                    m.start -= 1
        found.append((p, matches))
    engine.finalize(include_uncited=include_uncited)

    # pass 2: replace, last match first so earlier offsets stay valid
    report = ProcessReport(style=engine.style.name, output=os.fspath(output_path))
    for p, matches in found:
        rendered = [(m, engine.format_cluster(m.cluster)) for m in matches]
        for m, rich in reversed(rendered):
            replace_span(p, m.start, m.end, rich)
            report.clusters += 1

    # bibliography
    entries = engine.bibliography()
    doc_locale = locale if locale in ("en", "fa") else _detect_doc_locale(paragraphs)
    title = engine.bibliography_title(doc_locale) if bib_title is None else bib_title
    if bib_anchor is not None:
        report.placeholder_found = True
        template_r = _clean_template_run(bib_anchor)
        last = bib_anchor
        if bib_title:   # explicit heading requested
            last = _insert_paragraph_after(last)
            last.add_run(bib_title).bold = True
        for e in entries:
            last = _insert_paragraph_after(last)
            _fill_entry(last, e, engine.style, template_r)
        bib_anchor._p.getparent().remove(bib_anchor._p)
    elif entries:
        if title:
            try:
                h = doc.add_heading(title, level=1)
            except KeyError:
                h = doc.add_paragraph()
                h.add_run(title).bold = True
            if doc_locale == "fa":
                _set_bidi(h)
        for e in entries:
            _fill_entry(doc.add_paragraph(), e, engine.style)

    doc.save(os.fspath(output_path))
    report.cited = list(engine.cited)
    report.missing = list(engine.missing)
    report.uncited = [] if include_uncited else engine.uncited
    report.bibliography_entries = len(entries)
    return report
