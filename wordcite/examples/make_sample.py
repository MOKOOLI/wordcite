"""Create examples/sample.docx: an English + Persian draft with citation markers."""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).parent


def rtl(paragraph):
    paragraph._p.get_or_add_pPr().insert(0, OxmlElement("w:bidi"))
    for r in paragraph.runs:
        r.font.rtl = True


def main(path=HERE / "sample.docx"):
    doc = Document()
    doc.add_heading("Early sepsis detection: a short review", level=1)
    p = doc.add_paragraph("Machine learning can flag sepsis hours earlier than clinical scores [@smith2020]. ")
    # a marker split across runs with mixed formatting, like Word often produces
    p.add_run("Readmission risk follows a similar pattern [@smith2020b").italic = True
    p.add_run("; @garcia2021, p. 1045].")
    doc.add_paragraph("As @kahneman2011 argued, fast intuitions are often wrong, "
                      "and ethics frameworks matter [see @brown2019; @who2023].")
    doc.add_paragraph("Privacy-preserving approaches are maturing [@lee2022; @nguyen2018]. "
                      "Contact: someone@example.com (this is not a citation).")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Federated learning"
    t.rows[0].cells[1].text = "Reported in [-@lee2022] and \\cite{garcia2021}."
    rtl(doc.add_paragraph("یادگیری عمیق در تشخیص پزشکی کاربرد گسترده‌ای دارد [@rezaei1399; @karimi1398]."))
    doc.add_heading("References", level=1)
    doc.add_paragraph("[[BIBLIOGRAPHY]]")
    doc.save(path)
    return path


if __name__ == "__main__":
    print(main())
