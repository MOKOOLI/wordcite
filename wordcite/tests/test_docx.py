from pathlib import Path

from docx import Document

from wordcite import process_docx

ROOT = Path(__file__).resolve().parents[1]
BIB = ROOT / "examples" / "references.bib"


def _doc(tmp_path, paragraphs, split=None):
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    if split:
        p = d.add_paragraph()
        for chunk in split:
            p.add_run(chunk)
    path = tmp_path / "in.docx"
    d.save(path)
    return path


def _texts(path):
    return [p.text for p in Document(path).paragraphs]


def test_end_to_end_apa(tmp_path):
    src = _doc(tmp_path, ["Intro [@smith2020].", "As @kahneman2011 said.", "[[BIBLIOGRAPHY]]"])
    out = tmp_path / "out.docx"
    rep = process_docx(src, out, BIB, style="apa")
    texts = _texts(out)
    assert texts[0] == "Intro (Smith & Jones, 2020)."
    assert texts[1] == "As Kahneman (2011) said."
    assert texts[2].startswith("Kahneman, D. (2011).")
    assert len(texts) == 4 and rep.clusters == 2 and rep.placeholder_found


def test_marker_split_across_runs(tmp_path):
    src = _doc(tmp_path, [], split=["Text [@smi", "th2020; @garcia2021", ", p. 4] end."])
    out = tmp_path / "out.docx"
    process_docx(src, out, BIB, style="vancouver")
    assert _texts(out)[0] == "Text [1,2, p. 4] end."


def test_superscript_removes_space(tmp_path):
    src = _doc(tmp_path, ["Known fact [@smith2020]."])
    out = tmp_path / "out.docx"
    process_docx(src, out, BIB, style="nature")
    p = Document(out).paragraphs[0]
    assert p.text == "Known fact1."
    assert any(r.font.superscript and r.text == "1" for r in p.runs)


def test_appends_bibliography_without_placeholder(tmp_path):
    src = _doc(tmp_path, ["See [@who2023]."])
    out = tmp_path / "out.docx"
    rep = process_docx(src, out, BIB, style="ieee")
    texts = _texts(out)
    assert texts[0] == "See [1]."
    assert texts[1] == "References"
    assert texts[2].startswith("[1]\tWorld Health Organization.")
    assert not rep.placeholder_found


def test_emails_untouched_and_missing_reported(tmp_path):
    src = _doc(tmp_path, ["Mail me at a@b.com or @nobody [@ghost]."])
    out = tmp_path / "out.docx"
    rep = process_docx(src, out, BIB, style="apa")
    assert _texts(out)[0].startswith("Mail me at a@b.com or @nobody (ghost?)")
    assert rep.missing == ["ghost"]
